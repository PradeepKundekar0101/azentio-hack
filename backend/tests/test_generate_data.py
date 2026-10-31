"""Synthetic data generator: determinism, identical exports, and an answer key verified by an independent checker."""

import copy
import json
import re
from collections import Counter
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.core.money import parse_amount
from scripts.datagen.builder import PLAN_COUNTS
from scripts.datagen.common import DATASETS, parse_date
from scripts.datagen.oracle import verify
from scripts.datagen.readback import read_all, read_csv_dir
from scripts.datagen.render import HEADERS, SHEETS
from scripts.generate_data import DEFAULT_OUT, build, generate

SMALL = 0.05
ERROR_CODES = {
    "missing_field", "negative_amount", "invalid_date", "ambiguous_date", "submission_before_incident",
    "unknown_policy", "unknown_claim", "line_amount_mismatch", "duplicate_id_in_file", "unknown_service_code",
}  # fmt: skip

# The importer's header synonyms (compact form), from the Phase 2 contract.
SYNONYMS = {
    "claims": {
        "claimid": {"claimno", "claimnumber"}, "policynumber": {"policyno", "policynum"},
        "claimantid": {"memberid", "insuredid"}, "claimantname": {"membername", "patientname"},
        "incidentdate": {"dateofincident", "admissiondate"}, "submissiondate": {"dateofsubmission", "submitteddate"},
        "claimtype": set(), "requestedamount": {"claimamount", "amountclaimed", "claimedamount"},
        "incidentreference": {"incidentref"}, "status": {"claimstatus"},
    },
    "line_items": {
        "lineid": {"lineno", "lineitemid"}, "claimid": set(),
        "invoicereference": {"invoiceno", "invoicenumber", "invoiceref", "billno"},
        "providerid": {"providercode", "hospitalid"}, "providername": {"hospitalname"},
        "servicedate": {"dateofservice"}, "servicecode": {"procedurecode"}, "servicedescription": {"description"},
        "category": {"servicecategory"}, "quantity": {"qty", "units"}, "unitprice": {"rate", "unitrate"},
        "lineamount": {"amount", "linetotal"},
    },
    "policies": {
        "policynumber": {"policyno"}, "policyholdername": {"policyholder"},
        "insuredmemberids": {"memberids", "insuredmembers"}, "productname": {"product", "plan"},
        "benefitperiodstart": {"periodstart", "coverstart"}, "benefitperiodend": {"periodend", "coverend"},
        "overalllimit": {"suminsured"}, "deductibleamount": {"deductible"}, "deductiblebasis": {"deductibletype"},
        "copaypercent": {"copay", "copaypct"},
    },
    "policy_rules": {
        "ruleid": {"ruleno"}, "policynumber": set(), "ruletype": set(), "scopetype": set(), "scopevalue": set(),
        "limitbasis": set(), "limitamount": set(), "clauseref": {"clause", "clausereference"}, "description": set(),
    },
    "settlements": {
        "settlementid": {"settlementno"}, "claimid": {"claimno"}, "settlementdate": set(),
        "approvedamount": {"approved"}, "paidamount": {"paid", "amountpaid"}, "settlementstatus": {"status"},
    },
}  # fmt: skip
SHEET_SYNONYMS = {
    "policies": {"policies"},
    "policy_rules": {"policyrules", "rules"},
    "claims": {"claims"},
    "line_items": {"lineitems", "lines", "claimlines"},
    "settlements": {"settlements", "payments"},
}


def compact(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", text.lower())


@pytest.fixture(scope="module")
def small(tmp_path_factory: pytest.TempPathFactory) -> Path:
    out = tmp_path_factory.mktemp("sample")
    generate(out, seed=42, scale=SMALL)
    return out


@pytest.fixture(scope="module")
def full() -> tuple[dict, dict]:
    _, rows, key = build(42, 1.0, demo=False)
    return rows, key


def tree(directory: Path) -> dict[str, bytes]:
    return {str(p.relative_to(directory)): p.read_bytes() for p in sorted(directory.rglob("*")) if p.is_file()}


def test_same_seed_produces_identical_files(small: Path, tmp_path: Path) -> None:
    generate(tmp_path / "again", seed=42, scale=SMALL)
    first, second = tree(small), tree(tmp_path / "again")
    assert sorted(first) == sorted(second)
    assert len(first) == 16  # 5 csv + xlsx + json + key, for the full set and the demo pack
    assert first == second


def test_another_seed_produces_different_data(small: Path, tmp_path: Path) -> None:
    generate(tmp_path / "other", seed=7, scale=SMALL)
    assert tree(small)["csv/claims.csv"] != tree(tmp_path / "other")["csv/claims.csv"]


@pytest.mark.parametrize("pack", ["", "demo"])
def test_csv_excel_and_json_contain_the_same_records(small: Path, pack: str) -> None:
    formats = read_all(small / pack)
    assert formats["csv"] == formats["xlsx"] == formats["json"]
    key = json.loads((small / pack / "answer_key.json").read_text())
    assert {ds: len(rows) for ds, rows in formats["csv"].items()} == key["counts"]


@pytest.mark.parametrize("pack", ["", "demo"])
def test_answer_key_is_verified_by_the_independent_checker(small: Path, pack: str) -> None:
    rows = read_csv_dir(small / pack / "csv")
    key = json.loads((small / pack / "answer_key.json").read_text())
    assert verify(rows, key) == []


def test_checker_catches_a_wrong_answer_key_and_a_tampered_row(small: Path) -> None:
    rows = read_csv_dir(small / "csv")
    key = json.loads((small / "answer_key.json").read_text())
    wrong = copy.deepcopy(key)
    wrong["findings"][0]["expected_amount"] = "1.00"
    assert any("not recomputed" in p for p in verify(rows, wrong))
    tampered = copy.deepcopy(rows)
    victim = next(r for r in tampered["line_items"] if r["service_code"] == "CON-GP")
    victim["line_amount"] = "999999"
    assert verify(tampered, key) != []


def test_full_scale_matches_the_plan(full: tuple[dict, dict]) -> None:
    rows, key = full
    assert verify(rows, key) == []
    scenarios = Counter(f["scenario"] for f in key["findings"])
    expected = {s: n for s, n in PLAN_COUNTS.items() if s != "legitimate_repeat"}
    assert dict(scenarios) == expected
    assert len(key["legitimate_repeats"]) == PLAN_COUNTS["legitimate_repeat"]
    accepted = key["accepted_counts"]
    assert accepted["policies"] == 300
    claims = [r for r in rows["claims"] if r["status"] in ("open", "settled")]
    statuses = Counter(r["status"] for r in claims)
    assert 1450 <= statuses["open"] <= 1650 and 1450 <= statuses["settled"] <= 1550
    assert 11000 <= accepted["line_items"] <= 14000
    total = sum(key["counts"].values())
    rejected = sum(1 for i in key["invalid_rows"] if i["severity"] == "error")
    assert 0.01 <= rejected / total <= 0.03
    codes = Counter(i["error_code"] for i in key["invalid_rows"])
    for code in ERROR_CODES:
        assert codes[code] >= 5, code
    assert codes["requested_amount_mismatch"] >= 5
    assert set(codes) <= ERROR_CODES | {"requested_amount_mismatch", "parent_rejected"}


def test_finding_entries_have_the_contract_shape(full: tuple[dict, dict]) -> None:
    _, key = full
    checks = {
        "claimant_not_insured": ("NOT_COVERED", {"C02"}),
        "not_covered": ("NOT_COVERED", {"C03", "C04"}),
        "outside_benefit_period": ("OUT_OF_BENEFIT_PERIOD", {"C01"}),
        "excluded_service": ("EXCLUDED_SERVICE", {"C05"}),
        "room_rent_above_per_day_limit": ("SUB_LIMIT_EXCEEDED", {"C06"}),
        "per_claim_sub_limit": ("SUB_LIMIT_EXCEEDED", {"C07"}),
        "per_period_sub_limit": ("SUB_LIMIT_EXCEEDED", {"C08"}),
        "overall_limit": ("OVERALL_LIMIT_EXCEEDED", {"C11"}),
        "settled_without_copay": ("SETTLEMENT_OVERPAYMENT", {"C12"}),
        "settled_without_deductible": ("SETTLEMENT_OVERPAYMENT", {"C12"}),
        "settled_above_approved": ("SETTLEMENT_OVERPAYMENT", {"C12"}),
        "exact_duplicate_claim": ("EXACT_DUPLICATE_CLAIM", {None}),
        "duplicate_line_within_claim": ("DUPLICATE_LINE_ITEM", {None}),
        "duplicate_line_across_claims": ("DUPLICATE_LINE_ITEM", {None}),
        "near_duplicate": ("NEAR_DUPLICATE", {None}),
    }
    for finding in key["findings"]:
        category, codes = checks[finding["scenario"]]
        assert finding["category"] == category
        assert finding["check_code"] in codes
        assert finding["amount_type"] == ("estimated" if finding["scenario"] == "near_duplicate" else "calculated")
        assert Decimal(finding["expected_amount"]) > 0
        assert re.fullmatch(r"\d+\.\d{2}", finding["expected_amount"])
        if finding["scenario"] in ("exact_duplicate_claim", "near_duplicate", "duplicate_line_across_claims"):
            assert finding["matched_claim_id"] and finding["matched_claim_id"] != finding["claim_id"]
        if finding["check_code"] in ("C04", "C05", "C06") or finding["scenario"].startswith("duplicate_line"):
            assert finding["line_id"]
        if finding["check_code"] in ("C05", "C06", "C07", "C08"):
            assert finding["policy_rule_id"]


def test_valid_line_items_sum_to_their_claims_requested_amount(full: tuple[dict, dict]) -> None:
    rows, key = full
    rejected = {(i["dataset"], i["row_index"]) for i in key["invalid_rows"] if i["severity"] == "error"}
    mismatches = {i["record_id"] for i in key["invalid_rows"] if i["error_code"] == "requested_amount_mismatch"}
    totals: dict[str, Decimal] = Counter()  # type: ignore[assignment]
    for index, row in enumerate(rows["line_items"]):
        if ("line_items", index) not in rejected:
            totals[row["claim_id"]] += parse_amount(row["line_amount"])
    checked = 0
    for index, row in enumerate(rows["claims"]):
        if ("claims", index) in rejected:
            continue
        differs = abs(totals[row["claim_id"]] - parse_amount(row["requested_amount"])) > 1
        assert differs == (row["claim_id"] in mismatches), row["claim_id"]
        checked += 1
    assert checked > 2500


def test_dates_and_amounts_use_only_accepted_formats_outside_planted_rows(full: tuple[dict, dict]) -> None:
    rows, key = full
    planted = {(i["dataset"], i["row_index"]) for i in key["invalid_rows"]}
    date_fields = {"benefit_period_start", "benefit_period_end", "incident_date", "submission_date", "service_date"}
    date_fields.add("settlement_date")
    money_fields = {"overall_limit", "deductible_amount", "limit_amount", "requested_amount", "unit_price"}
    money_fields |= {"line_amount", "approved_amount", "paid_amount"}
    styles = Counter()
    for dataset in DATASETS:
        for index, row in enumerate(rows[dataset]):
            if (dataset, index) in planted:
                continue
            for name, value in row.items():
                if value is None:
                    continue
                if name in date_fields:
                    parsed, error = parse_date(value)
                    assert error is None and parsed is not None, (dataset, index, name, value)
                    styles["iso" if value[4] == "-" else "dmy"] += 1
                if name in money_fields:
                    assert parse_amount(value) >= 0
                    styles["rupee" if "₹" in value else "grouped" if "," in value else "plain"] += 1
    assert all(styles[s] > 1000 for s in ("iso", "dmy", "rupee", "grouped", "plain"))


def test_headers_and_sheet_names_match_the_importer_synonyms() -> None:
    for fmt, datasets in HEADERS.items():
        for dataset, headers in datasets.items():
            allowed = {name: {name} | synonyms for name, synonyms in SYNONYMS[dataset].items()}
            assert [next(n for n, names in allowed.items() if compact(h) in names) for h in headers] == list(allowed), (
                fmt,
                dataset,
            )
    for dataset, sheet in SHEETS.items():
        assert compact(sheet) in SHEET_SYNONYMS[dataset]
    assert len({tuple(HEADERS[fmt]["claims"]) for fmt in HEADERS}) == 3


def test_excel_cells_are_text(small: Path) -> None:
    workbook = load_workbook(small / "xlsx" / "claims_data.xlsx", read_only=True)
    assert [compact(name) for name in workbook.sheetnames] == [
        "policies",
        "policyrules",
        "claims",
        "lineitems",
        "settlements",
    ]
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            assert all(value is None or isinstance(value, str) for value in row)
    workbook.close()


def test_json_members_are_arrays_and_csv_members_are_pipe_separated(small: Path) -> None:
    payload = json.loads((small / "json" / "claims_data.json").read_text())
    assert sorted(payload) == sorted(DATASETS)
    assert all(isinstance(p["insured_member_ids"], list) for p in payload["policies"])
    csv_text = (small / "csv" / "policies.csv").read_text()
    assert re.search(r"MEM-\d{6}\|MEM-\d{6}", csv_text)


def test_demo_pack_is_small_self_contained_and_covers_every_scenario(small: Path) -> None:
    demo = small / "demo"
    rows = read_csv_dir(demo / "csv")
    key = json.loads((demo / "answer_key.json").read_text())
    assert verify(rows, key) == []
    assert 45 <= key["accepted_counts"]["claims"] <= 60
    scenarios = Counter(f["scenario"] for f in key["findings"])
    assert set(scenarios) == set(PLAN_COUNTS) - {"legitimate_repeat"}
    assert min(scenarios.values()) >= 2
    assert len(key["legitimate_repeats"]) >= 2
    assert {"negative_amount", "ambiguous_date", "unknown_claim"} <= {i["error_code"] for i in key["invalid_rows"]}
    policies = {r["policy_number"] for r in rows["policies"]}
    unknown = {i["record_id"] for i in key["invalid_rows"] if i["error_code"] == "unknown_policy"}
    assert all(r["policy_number"] in policies or r["claim_id"] in unknown for r in rows["claims"])
    amounts = {(f["scenario"], f["expected_amount"]) for f in key["findings"]}
    assert ("room_rent_above_per_day_limit", "6000.00") in amounts  # (7,000 - 5,000) x 3
    assert ("settled_without_copay", "8400.00") in amounts  # paid 84,000; entitlement 75,600
    refs = {r["invoice_reference"] for r in rows["line_items"]}
    assert {"INV-00452", "INV452A"} <= refs


def test_committed_sample_is_current(tmp_path: Path) -> None:
    if not (DEFAULT_OUT / "answer_key.json").exists():
        pytest.skip("data/sample has not been generated")
    generate(tmp_path, seed=42, scale=1.0)
    committed = {k: v for k, v in tree(DEFAULT_OUT).items() if not k.startswith("invoices/")}
    assert committed == tree(tmp_path)
