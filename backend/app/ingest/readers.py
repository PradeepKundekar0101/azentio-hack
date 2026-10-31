"""Read CSV, Excel and JSON uploads into tables of raw text rows, keeping file, sheet and row number."""

import csv
import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import PurePath
from typing import Any

from openpyxl import load_workbook

from app.ingest.mapping import Dataset, dataset_from_headers, dataset_from_name

CSV_DELIMITERS = ",;\t"


class ReadError(ValueError):
    """The file as a whole cannot be read."""


@dataclass
class RawRow:
    row_number: int
    values: dict[str, Any]  # original header -> raw value (text, or a JSON value)


@dataclass
class RawTable:
    file_name: str
    sheet: str | None
    dataset: Dataset | None
    headers: list[str]
    rows: list[RawRow] = field(default_factory=list)


def read_upload(file_name: str, data: bytes, dataset_hint: Dataset | None = None) -> list[RawTable]:
    suffix = PurePath(file_name).suffix.lower()
    if suffix == ".csv":
        return [_read_csv(file_name, data, dataset_hint)]
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(file_name, data)
    if suffix == ".json":
        return _read_json(file_name, data, dataset_hint)
    raise ReadError(f"Unsupported file type {suffix or '(none)'}; upload CSV, Excel (.xlsx) or JSON")


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ReadError("File is not UTF-8 or Latin-1 text")  # pragma: no cover - latin-1 decodes any bytes


def _read_csv(file_name: str, data: bytes, dataset_hint: Dataset | None) -> RawTable:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=CSV_DELIMITERS)
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        raise ReadError("File is empty")
    headers = [h.strip() for h in rows[0]]
    dataset = dataset_hint or dataset_from_name(PurePath(file_name).stem) or dataset_from_headers(headers)
    table = RawTable(file_name, None, dataset, headers)
    for number, row in enumerate(rows[1:], start=2):
        if not any(cell.strip() for cell in row):
            continue
        padded = row + [""] * (len(headers) - len(row))
        table.rows.append(RawRow(number, dict(zip(headers, padded, strict=False))))
    return table


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, Decimal | int):
        return str(value)
    return str(value)


def _read_xlsx(file_name: str, data: bytes) -> list[RawTable]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ReadError("File is not a readable Excel workbook") from exc
    tables: list[RawTable] = []
    try:
        for sheet in workbook.worksheets:
            rows = [[_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            if not rows or not any(rows[0]):
                continue
            headers = [h.strip() for h in rows[0]]
            while headers and not headers[-1]:
                headers.pop()
            dataset = dataset_from_name(sheet.title) or dataset_from_headers(headers)
            table = RawTable(file_name, sheet.title, dataset, headers)
            for number, row in enumerate(rows[1:], start=2):
                cells = row[: len(headers)]
                if not any(c.strip() for c in cells):
                    continue
                cells = cells + [""] * (len(headers) - len(cells))
                table.rows.append(RawRow(number, dict(zip(headers, cells, strict=True))))
            tables.append(table)
    finally:
        workbook.close()
    return tables


def _json_table(file_name: str, dataset: Dataset | None, key: str | None, records: Any) -> RawTable:
    if not isinstance(records, list) or not all(isinstance(r, dict) for r in records):
        raise ReadError(f"{key or 'The JSON array'} must be a list of objects")
    headers: list[str] = []
    for record in records:
        for header in record:
            if header not in headers:
                headers.append(header)
    table = RawTable(file_name, key, dataset, headers)
    for index, record in enumerate(records, start=1):
        table.rows.append(RawRow(index, {h: record.get(h) for h in headers}))
    return table


def _read_json(file_name: str, data: bytes, dataset_hint: Dataset | None) -> list[RawTable]:
    try:
        document = json.loads(_decode(data))
    except json.JSONDecodeError as exc:
        raise ReadError(f"Invalid JSON: {exc.msg} at line {exc.lineno}") from exc
    if isinstance(document, list):
        dataset = dataset_hint or dataset_from_name(PurePath(file_name).stem)
        table = _json_table(file_name, dataset, None, document)
        if table.dataset is None:
            table.dataset = dataset_from_headers(table.headers)
        return [table]
    if isinstance(document, dict):
        return [_json_table(file_name, dataset_from_name(key), key, records) for key, records in document.items()]
    raise ReadError("JSON must be an object with one key per dataset, or an array of records")
