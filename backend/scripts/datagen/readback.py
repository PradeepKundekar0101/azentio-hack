"""Reads the exported CSV, Excel and JSON files back into canonical rows (field name -> raw text)."""

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from scripts.datagen.common import DATASETS
from scripts.datagen.render import COLUMNS, HEADERS, SHEETS, Row


def _canonical(fmt: str, dataset: str, raw: dict[str, Any]) -> Row:
    mapping = dict(zip(HEADERS[fmt][dataset], COLUMNS[dataset], strict=True))
    row: Row = {}
    for header, value in raw.items():
        field = mapping[header]
        if value is None or value == "":
            row[field] = None
        elif field == "insured_member_ids":
            row[field] = list(value) if isinstance(value, list) else str(value).split("|")
        elif isinstance(value, bool):
            raise ValueError(f"unexpected boolean in {dataset}.{field}")
        elif isinstance(value, float):
            row[field] = repr(value)
        else:
            row[field] = str(value)
    return row


def read_csv_dir(directory: Path) -> dict[str, list[Row]]:
    out: dict[str, list[Row]] = {}
    for dataset in DATASETS:
        with (directory / f"{dataset}.csv").open(newline="", encoding="utf-8") as handle:
            out[dataset] = [_canonical("csv", dataset, raw) for raw in csv.DictReader(handle)]
    return out


def read_xlsx(path: Path) -> dict[str, list[Row]]:
    workbook = load_workbook(path, read_only=True)
    out: dict[str, list[Row]] = {}
    for dataset in DATASETS:
        rows = list(workbook[SHEETS[dataset]].iter_rows(values_only=True))
        headers = [str(h) for h in rows[0]]
        width = len(headers)
        out[dataset] = [
            _canonical("xlsx", dataset, dict(zip(headers, (*values, *[None] * (width - len(values))), strict=True)))
            for values in rows[1:]
        ]
    workbook.close()
    return out


def read_json(path: Path) -> dict[str, list[Row]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {dataset: [_canonical("json", dataset, raw) for raw in payload[dataset]] for dataset in DATASETS}


def read_all(directory: Path) -> dict[str, dict[str, list[Row]]]:
    return {
        "csv": read_csv_dir(directory / "csv"),
        "xlsx": read_xlsx(directory / "xlsx" / "claims_data.xlsx"),
        "json": read_json(directory / "json" / "claims_data.json"),
    }
